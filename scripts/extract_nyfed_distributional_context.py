#!/usr/bin/env python3
"""Extract exact-bound New York Fed age distribution context from the retained Q2 2026 workbook.

This script is non-authorizing. It refuses a workbook hash or layout change and emits
contextual distribution evidence only. Debt stock is never relabeled as new borrowing,
and delinquency transitions are not promoted to a complete household-stress finding.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
MAP_PATH = ROOT / "research-data/household-economic-conditions/nyfed-2026q2-distributional-map.v1.json"
GOAL = "ERL-HOUSEHOLD-ECONOMIC-CONDITIONS-SITE-001"
PERIOD_RE = re.compile(r"^(\d{2}):Q([1-4])$")


def period(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    m = PERIOD_RE.fullmatch(value.strip())
    return f"20{m.group(1)}-Q{m.group(2)}" if m else None


def extract(raw: bytes, mapping: dict[str, Any]) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb[mapping["worksheet"]]
    header_row = int(mapping["header_row"])
    for col, expected in mapping["expected_headers"].items():
        observed = ws.cell(header_row, column_index_from_string(col)).value
        if observed != expected:
            raise ValueError(f"header mismatch {mapping['worksheet']}!{col}{header_row}: {observed!r} != {expected!r}")
    pcol = column_index_from_string(mapping["period_column"])
    rows = []
    for rowno in range(header_row + 1, ws.max_row + 1):
        p = period(ws.cell(rowno, pcol).value)
        if p is None:
            continue
        for category, col in mapping["category_columns"].items():
            value = ws.cell(rowno, column_index_from_string(col)).value
            if value in (None, ""):
                continue
            if not isinstance(value, (int, float)):
                raise ValueError(f"nonnumeric value at {mapping['worksheet']}!{col}{rowno}")
            rows.append({"period": p, "category": category, "value": float(value), "evidence_class": "DIRECT_OBSERVATION"})
    if not rows:
        raise ValueError(f"no observations for {mapping['worksheet']}")
    return rows


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--output", type=Path, required=True)
    args = ap.parse_args()
    raw = args.workbook.read_bytes()
    spec = json.loads(MAP_PATH.read_text(encoding="utf-8"))
    if spec["goal_task_id"] != GOAL or spec["finding_authority"] is not False or spec["public_activation_authorized"] is not False:
        raise SystemExit("distribution map authority boundary invalid")
    observed_hash = hashlib.sha256(raw).hexdigest()
    if observed_hash != spec["workbook_sha256"]:
        raise SystemExit(f"NY Fed distribution workbook hash mismatch: {observed_hash}")
    series = []
    for series_id, mapping in spec["maps"].items():
        observations = extract(raw, mapping)
        series.append({
            "series_id": series_id,
            "label": mapping["title"],
            "source_agency": "Federal Reserve Bank of New York Consumer Credit Panel/Equifax",
            "frequency": "QUARTERLY",
            "unit": mapping["unit"],
            "earliest_comparable_date": observations[0]["period"],
            "vintage": "2026Q2",
            "revision_timestamp": None,
            "comparison_mode": "ABSOLUTE",
            "same_axis_group": series_id,
            "observations": observations,
            "structural_breaks": [],
            "source_url": spec["source_workbook"],
            "methodology_url": "https://www.newyorkfed.org/microeconomics/hhdc",
            "admission_role": mapping["admission_role"],
        })
    output = {
        "schema": "stegverse.erl.nyfed-distributional-context/v1",
        "goal_task_id": GOAL,
        "release_id": spec["release_id"],
        "raw_sha256": observed_hash,
        "series": series,
        "semantic_boundaries": spec["semantic_boundaries"],
        "finding_authority": False,
        "public_activation_authorized": False,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print("NYFED_DISTRIBUTIONAL_CONTEXT=PASS")
    print("DEBT_STOCK_PROMOTED_TO_NEW_BORROWING=false")
    print("PUBLIC_ACTIVATION_AUTHORIZED=false")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
