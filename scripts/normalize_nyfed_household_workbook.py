#!/usr/bin/env python3
"""Normalize only the exact, observed NY Fed Q2 2026 workbook surfaces.

This module is non-authorizing. It validates worksheet/header/column identity before
emitting observations, so a changed workbook layout fails closed rather than being
silently reinterpreted.
"""
from __future__ import annotations

import argparse
import io
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
BINDINGS_PATH = ROOT / "research-data/household-economic-conditions/official-series-bindings.v1.json"
GOAL = "ERL-HOUSEHOLD-ECONOMIC-CONDITIONS-SITE-001"
PERIOD_RE = re.compile(r"^(\d{2}):Q([1-4])$")


def canonical_bytes(value: Any) -> bytes:
    return (json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False) + "\n").encode("utf-8")


def load_bindings() -> dict[str, dict[str, Any]]:
    document = json.loads(BINDINGS_PATH.read_text(encoding="utf-8"))
    return {row["inventory_series_id"]: row for row in document["bindings"]}


def normalize_period(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    match = PERIOD_RE.fullmatch(value.strip())
    if not match:
        return None
    return f"20{match.group(1)}-Q{match.group(2)}"


def normalize_workbook(raw: bytes, binding: dict[str, Any]) -> list[dict[str, Any]]:
    mapping = binding["normalization"]
    if mapping.get("parser_state") != "EXACT_WORKBOOK_MAP_BOUND":
        raise ValueError("NY Fed parser map is not exact-bound")

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet_name = mapping["worksheet"]
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"NY Fed workbook missing worksheet {sheet_name!r}")
    worksheet = workbook[sheet_name]

    header_row = int(mapping["header_row"])
    period_column = mapping["period_column"]
    expected_headers = mapping["expected_headers"]
    for column, expected in expected_headers.items():
        observed = worksheet.cell(header_row, column_index_from_string(column)).value
        if observed != expected:
            raise ValueError(
                f"NY Fed workbook header mismatch at {sheet_name}!{column}{header_row}: "
                f"expected {expected!r}, observed {observed!r}"
            )

    output = []
    period_index = column_index_from_string(period_column)
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        period = normalize_period(worksheet.cell(row_number, period_index).value)
        if period is None:
            continue
        for category, column in mapping["category_columns"].items():
            value = worksheet.cell(row_number, column_index_from_string(column)).value
            if value in (None, ""):
                continue
            if not isinstance(value, (int, float)):
                raise ValueError(f"NY Fed nonnumeric value at {sheet_name}!{column}{row_number}: {value!r}")
            output.append({
                "period": period,
                "category": category,
                "value": float(value),
                "evidence_class": "DIRECT_OBSERVATION",
            })
    if not output:
        raise ValueError(f"NY Fed worksheet {sheet_name!r} produced no observations")
    return output


def rewrite_candidate(candidate_path: Path, raw: bytes, binding: dict[str, Any]) -> int:
    candidate = json.loads(candidate_path.read_text(encoding="utf-8"))
    if candidate.get("goal_task_id") != GOAL:
        raise ValueError(f"unexpected goal task in {candidate_path}")
    observations = normalize_workbook(raw, binding)
    candidate["status"] = "NORMALIZED_SOURCE_OBSERVATIONS"
    candidate["observations"] = observations
    candidate["finding_authority"] = False
    candidate["public_activation_authorized"] = False
    candidate_path.write_bytes(canonical_bytes(candidate))
    return len(observations)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--candidate-root", type=Path, required=True)
    args = parser.parse_args()

    raw = args.workbook.read_bytes()
    bindings = load_bindings()
    ids = ["NYFED_CCP_DEBT_BALANCE_BY_CLASS", "NYFED_CCP_DELINQUENCY_BY_CLASS"]
    for series_id in ids:
        candidate_path = args.candidate_root / f"{series_id.lower()}.candidate.json"
        count = rewrite_candidate(candidate_path, raw, bindings[series_id])
        print(f"{series_id}=NORMALIZED_SOURCE_OBSERVATIONS observations={count}")
    print("NYFED_WORKBOOK_NORMALIZATION=EXACT_MAP_BOUND_PASS")
    print("PUBLIC_ACTIVATION_AUTHORIZED=false")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
