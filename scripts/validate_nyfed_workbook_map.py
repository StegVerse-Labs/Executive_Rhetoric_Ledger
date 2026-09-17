#!/usr/bin/env python3
"""Deterministic conformance tests for the exact NY Fed Q2 2026 workbook map."""
from __future__ import annotations

import importlib.util
import io
import json
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
MAP_PATH = ROOT / "research-data/household-economic-conditions/nyfed-2026q2-workbook-map.v1.json"
NORMALIZER = ROOT / "scripts/normalize_nyfed_household_workbook.py"
GOAL = "ERL-HOUSEHOLD-ECONOMIC-CONDITIONS-SITE-001"

spec = importlib.util.spec_from_file_location("nyfed_normalizer", NORMALIZER)
module = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(module)

mapping_document = json.loads(MAP_PATH.read_text(encoding="utf-8"))
assert mapping_document["goal_task_id"] == GOAL
assert mapping_document["release_id"] == "2026Q2"
assert mapping_document["workbook_sha256"] == "ddfba16b87e187848ed591a1283f310188e7bdeebe5cf409b70a35b058d99237"
assert mapping_document["finding_authority"] is False
assert mapping_document["public_activation_authorized"] is False

maps = mapping_document["maps"]
debt = maps["NYFED_CCP_DEBT_BALANCE_BY_CLASS"]
delinquency = maps["NYFED_CCP_DELINQUENCY_BY_CLASS"]
assert debt["worksheet"] == "Page 3 Data" and debt["header_row"] == 4 and debt["period_column"] == "A"
assert debt["category_columns"] == {"Mortgage":"B","HELOC":"C","Auto Loan":"D","Credit Card":"E","Student Loan":"F","Other":"G","Total":"H"}
assert delinquency["worksheet"] == "Page 14 Data" and delinquency["header_row"] == 5 and delinquency["period_column"] == "A"
assert delinquency["category_columns"] == {"Auto Loan":"B","Credit Card":"C","Mortgage":"D","HELOC":"E","Student Loan":"F","Other":"G","ALL":"H"}

wb = Workbook()
ws = wb.active
ws.title = "Page 3 Data"
for column, value in debt["expected_headers"].items():
    ws[f"{column}4"] = value
ws["A5"] = "26:Q1"
ws["A6"] = "26:Q2"
for row, base in ((5, 1.0), (6, 2.0)):
    for offset, column in enumerate("BCDEFGH", start=1):
        ws[f"{column}{row}"] = base + offset / 10

ws2 = wb.create_sheet("Page 14 Data")
for column, value in delinquency["expected_headers"].items():
    ws2[f"{column}5"] = value
ws2["A6"] = "26:Q1"
ws2["A7"] = "26:Q2"
for row, base in ((6, 3.0), (7, 4.0)):
    for offset, column in enumerate("BCDEFGH", start=1):
        ws2[f"{column}{row}"] = base + offset / 10

buf = io.BytesIO()
wb.save(buf)
raw = buf.getvalue()

debt_rows = module.normalize_workbook(raw, debt)
delinquency_rows = module.normalize_workbook(raw, delinquency)
assert len(debt_rows) == 14
assert len(delinquency_rows) == 14
assert debt_rows[0] == {"period":"2026-Q1","category":"Mortgage","value":1.1,"evidence_class":"DIRECT_OBSERVATION"}
assert debt_rows[-1] == {"period":"2026-Q2","category":"Total","value":2.7,"evidence_class":"DIRECT_OBSERVATION"}
assert delinquency_rows[0] == {"period":"2026-Q1","category":"Auto Loan","value":3.1,"evidence_class":"DIRECT_OBSERVATION"}
assert delinquency_rows[-1] == {"period":"2026-Q2","category":"ALL","value":4.7,"evidence_class":"DIRECT_OBSERVATION"}

wb_bad = Workbook()
ws_bad = wb_bad.active
ws_bad.title = "Page 3 Data"
for column, value in debt["expected_headers"].items():
    ws_bad[f"{column}4"] = value
ws_bad["B4"] = "Changed Header"
ws_bad["A5"] = "26:Q2"
ws_bad["B5"] = 1.0
bad_buf = io.BytesIO()
wb_bad.save(bad_buf)
try:
    module.normalize_workbook(bad_buf.getvalue(), debt)
except ValueError:
    pass
else:
    raise AssertionError("changed NY Fed header did not fail closed")

print("NYFED_EXACT_WORKBOOK_MAP=PASS")
print("NYFED_DETERMINISTIC_NORMALIZATION=PASS")
print("NYFED_LAYOUT_CHANGE_FAIL_CLOSED=PASS")
print("PUBLIC_ACTIVATION_AUTHORIZED=false")
