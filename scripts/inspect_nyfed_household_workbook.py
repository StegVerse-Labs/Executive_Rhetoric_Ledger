#!/usr/bin/env python3
"""Inspect the official NY Fed Household Debt and Credit workbook structure.

This script records workbook structure and representative non-empty rows only. It does
not create findings or infer a parser map that has not been observed directly.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

GOAL = "ERL-HOUSEHOLD-ECONOMIC-CONDITIONS-SITE-001"


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, (int, float, bool, str)):
        return value
    return str(value)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--max-rows", type=int, default=18)
    parser.add_argument("--max-cols", type=int, default=16)
    args = parser.parse_args()

    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        nonempty = []
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, args.max_rows), values_only=True), start=1):
            values = [clean(v) for v in row[: args.max_cols]]
            if any(v not in (None, "") for v in values):
                nonempty.append({"row": r_idx, "values": values})
        sheets.append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "sample_nonempty_rows": nonempty,
        })

    observation = {
        "schema": "stegverse.erl.nyfed-household-workbook-structure-observation/v1",
        "goal_task_id": GOAL,
        "workbook_name": args.workbook.name,
        "workbook_sha256": sha256_file(args.workbook),
        "sheet_names": wb.sheetnames,
        "sheets": sheets,
        "finding_authority": False,
        "public_activation_authorized": False,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(observation, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(observation, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
